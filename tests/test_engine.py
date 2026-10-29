from io import BytesIO
import json
import pandas as pd
import pytest
from openpyxl import load_workbook
from prompt_compiler.demo import DEMOS, sample_files, demo_workflow
from prompt_compiler.files import inventory
from prompt_compiler.schema import Workflow
from prompt_compiler.validation import validate_workflow
from prompt_compiler.executor import execute, export_table
from prompt_compiler.compiler import compile_workflow
from prompt_compiler.provider import Reply, Provider
from prompt_compiler.benchmark import run_suite


def invoice_plan():
    return demo_workflow(DEMOS['Invoice review'], inventory(sample_files()))


def test_threshold_flag_preserves_rows_and_exports_cells():
    workflow = invoice_plan()
    run = execute(workflow, sample_files())
    assert run.error is None
    table = run.tables['report']
    assert len(table) == 4
    assert table['needs_review'].tolist() == [False, True, False, True]
    workbook = load_workbook(BytesIO(run.artifacts['invoice_report.xlsx']))
    rows = list(workbook.active.values)
    assert rows[0] == tuple(table.columns)
    assert [row[2] for row in rows[1:]] == [50000, 50001, 12500, 75000]
    assert [row[-1] for row in rows[1:]] == [False, True, False, True]


def test_filter_is_not_flag():
    files = sample_files()
    plan = demo_workflow(DEMOS['High-value invoices'], inventory(files))
    result = execute(plan, files)
    assert not result.error
    assert result.tables['report']['total'].tolist() == [50001, 75000]
    assert 'needs_review' not in result.tables['report']


def test_missing_total_preserved_as_unknown_flag():
    files = {'missing.txt': b'Vendor: Test\nCurrency: INR\n'}
    plan = demo_workflow(DEMOS['Invoice review'], inventory(files))
    result = execute(plan, files)
    assert not result.error
    assert pd.isna(result.tables['report'].iloc[0]['needs_review'])
    assert result.tables['report'].iloc[0]['extraction_status'] == 'needs_review'


def test_missing_filter_does_not_silently_drop_rows():
    files = {'missing.txt': b'Vendor: Test\nCurrency: INR\n'}
    plan = demo_workflow(DEMOS['High-value invoices'], inventory(files))
    result = execute(plan, files)
    assert 'missing values' in result.error
    assert result.artifacts == {}


@pytest.mark.parametrize('currency', ['USD', ''])
def test_currency_requires_explicit_inr(currency):
    files = {'foreign.txt': f'Vendor: Test\nCurrency: {currency}\nTotal: 60000\n'.encode()}
    plan = demo_workflow(DEMOS['Invoice review'], inventory(files))
    result = execute(plan, files)
    assert 'currency' in result.error


@pytest.mark.parametrize('mutation', ['cycle', 'duplicate', 'traversal', 'operation'])
def test_reject_invalid_structure(mutation):
    payload = invoice_plan().model_dump()
    if mutation == 'cycle':
        payload['steps'][1]['source'] = 'transform'
    if mutation == 'duplicate':
        payload['steps'][1]['id'] = 'read'
    if mutation == 'traversal':
        payload['steps'][-1]['filename'] = '../escape.xlsx'
    if mutation == 'operation':
        payload['steps'][1]['operation'] = 'shell'
    with pytest.raises(ValueError):
        Workflow.model_validate(payload)


def test_unknown_column_and_missing_file_rejected():
    manifest = inventory(sample_files())
    payload = invoice_plan().model_dump()
    payload['steps'][2]['column'] = 'invented'
    with pytest.raises(ValueError, match='unknown columns'):
        validate_workflow(Workflow.model_validate(payload), manifest)
    payload = invoice_plan().model_dump()
    payload['steps'][0]['files'] = ['not_uploaded.txt']
    with pytest.raises(ValueError, match='not uploaded'):
        validate_workflow(Workflow.model_validate(payload), manifest)


def test_formula_injection_escaped_in_workbook():
    table = pd.DataFrame({'=header': ['=HYPERLINK("https://example.invalid")', '  +cmd', '@formula', 'ordinary']})
    data = export_table(table, 'report.xlsx')
    worksheet = load_workbook(BytesIO(data), data_only=False).active
    assert worksheet['A1'].value == "'=header"
    assert worksheet['A2'].data_type != 'f'
    assert worksheet['A2'].value.startswith("'")
    assert worksheet['A5'].value == 'ordinary'


class StubProvider:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def complete(self, messages, schema=None):
        self.calls.append(messages.copy())
        return Reply(next(self.responses), 10, 5)


def envelope(workflow):
    return json.dumps({'workflow': workflow, 'clarification': None})


def test_feedback_repairs_invalid_column_and_tracks_tokens():
    good = invoice_plan().model_dump()
    bad = invoice_plan().model_dump()
    bad['steps'][2]['column'] = 'amount'
    provider = StubProvider(['read, extract, flag, export', envelope(bad), envelope(good)])
    result = compile_workflow(DEMOS['Invoice review'], inventory(sample_files()), provider)
    assert result.workflow is not None
    assert len(result.attempts) == 2
    assert result.attempts[0]['schema_valid']
    assert not result.attempts[0]['semantic_valid']
    assert result.prompt_tokens == 30 and result.completion_tokens == 15
    assert 'unknown columns' in provider.calls[-1][-1]['content']


def test_one_shot_does_not_retry():
    provider = StubProvider(['invalid json'])
    result = compile_workflow('task', inventory(sample_files()), provider, 'one_shot')
    assert result.workflow is None and len(result.attempts) == 1
    assert len(provider.calls) == 1


def test_clarification_not_retried_into_invented_workflow():
    provider = StubProvider([json.dumps({'workflow': None, 'clarification': 'Which exchange rate should be used?'})])
    result = compile_workflow('Convert currencies', inventory(sample_files()), provider, 'one_shot')
    assert 'Which exchange rate' in result.error
    assert result.workflow is None and len(provider.calls) == 1


def test_offline_custom_task_is_explicitly_unsupported():
    result = compile_workflow('Email my report', inventory(sample_files()))
    assert result.workflow is None
    assert 'exact sample instructions' in result.error


def test_development_benchmark_checks_reference_results():
    runs = run_suite()
    assert len(runs) == 3
    assert all(r['output_correct'] and r['workflow_correct'] for r in runs)
    assert all(r['mode'] == 'offline_demo' and r['manual_corrections'] is None for r in runs)


def test_csv_arithmetic_sort_aggregate_pipeline():
    files = {'sales.csv': b'vendor,quantity,price\nA,2,10\nB,3,8\nA,1,5\n'}
    plan = Workflow.model_validate({'version': '1.0', 'title': 'Sales', 'steps': [
        {'id': 'read', 'label': 'Read sales', 'operation': 'read', 'files': ['sales.csv']},
        {'id': 'calc', 'label': 'Calculate revenue', 'operation': 'calculate', 'source': 'read', 'left': 'quantity', 'right': 'price', 'arithmetic': 'multiply', 'output_column': 'revenue'},
        {'id': 'group', 'label': 'Group', 'operation': 'aggregate', 'source': 'calc', 'group_by': ['vendor'], 'column': 'revenue', 'function': 'sum', 'output_column': 'total_revenue'},
        {'id': 'sort', 'label': 'Sort', 'operation': 'sort', 'source': 'group', 'column': 'total_revenue', 'descending': True},
        {'id': 'export', 'label': 'Export', 'operation': 'export', 'source': 'sort', 'filename': 'sales.csv'}]})
    result = execute(plan, files)
    assert not result.error
    assert result.tables['export'].to_dict('records') == [{'vendor': 'A', 'total_revenue': 25}, {'vendor': 'B', 'total_revenue': 24}]


def test_provider_rejects_insecure_remote_and_credentials_in_url():
    with pytest.raises(ValueError):
        Provider('http://remote.example/v1', 'model')
    with pytest.raises(ValueError):
        Provider('https://user:secret@example.com/v1', 'model')


def test_feedback_budget_is_bounded():
    provider = StubProvider(['operation plan', 'invalid', 'invalid', 'invalid'])
    result = compile_workflow('task', inventory(sample_files()), provider, 'feedback', max_retries=2)
    assert result.workflow is None
    assert len(result.attempts) == 3 and len(provider.calls) == 4


def test_duplicate_exports_and_unused_nodes_rejected():
    payload = invoice_plan().model_dump()
    payload['steps'].append({**payload['steps'][-1], 'id': 'report_two'})
    with pytest.raises(ValueError, match='unique'):
        Workflow.model_validate(payload)
    payload = invoice_plan().model_dump()
    payload['steps'].insert(1, {'id': 'unused', 'label': 'Unused input', 'operation': 'read', 'files': ['invoice_001.txt']})
    with pytest.raises(ValueError, match='contribute'):
        Workflow.model_validate(payload)


def test_xlsx_input_roundtrip():
    data = export_table(pd.DataFrame({'vendor': ['A'], 'total': [12]}), 'input.xlsx')
    from prompt_compiler.files import read_file
    restored = read_file('input.xlsx', data)
    assert restored.to_dict('records') == [{'vendor': 'A', 'total': 12}]


def test_division_by_zero_and_invalid_numeric_input_fail():
    payload = {'version': '1.0', 'title': 'Calculation', 'steps': [
        {'id': 'read', 'label': 'Read', 'operation': 'read', 'files': ['test.csv']},
        {'id': 'calc', 'label': 'Divide', 'operation': 'calculate', 'source': 'read', 'left': 'value', 'right': 0.0, 'arithmetic': 'divide', 'output_column': 'result'},
        {'id': 'export', 'label': 'Export', 'operation': 'export', 'source': 'calc', 'filename': 'result.xlsx'}]}
    result = execute(Workflow.model_validate(payload), {'test.csv': b'value\n12\n'})
    assert 'Division by zero' in result.error and not result.artifacts
    payload['steps'][1]['right'] = 2.0
    result = execute(Workflow.model_validate(payload), {'test.csv': b'value\nnot_a_number\n'})
    assert 'non-numeric' in result.error and not result.artifacts


def test_real_http_adapter_against_local_stub():
    from http.server import BaseHTTPRequestHandler, HTTPServer
    import threading
    requests = []
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            requests.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
            response = json.dumps({'choices': [{'message': {'content': '{"ok":true}'}}],
                                   'usage': {'prompt_tokens': 12, 'completion_tokens': 4}}).encode()
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(response)
        def log_message(self, *args):
            pass
    server = HTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        provider = Provider(f'http://127.0.0.1:{server.server_port}/v1', 'test-model')
        reply = provider.complete([{'role': 'user', 'content': 'test'}], {'type': 'object'})
        assert reply.prompt_tokens == 12 and reply.completion_tokens == 4
        assert requests[0]['response_format']['json_schema']['strict']
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_strict_compilation_does_not_duplicate_schema_in_prompt():
    provider = StubProvider([envelope(invoice_plan().model_dump())])
    provider.structured = True
    result = compile_workflow(DEMOS['Invoice review'], inventory(sample_files()), provider, 'one_shot')
    assert result.workflow is not None
    assert 'JSON schema:' not in provider.calls[0][0]['content']
    provider = StubProvider([envelope(invoice_plan().model_dump())])
    provider.structured = False
    compile_workflow(DEMOS['Invoice review'], inventory(sample_files()), provider, 'one_shot')
    assert 'JSON schema:' in provider.calls[0][0]['content']


def test_presentation_escapes_uploaded_names_and_workflow_fields():
    from prompt_compiler.presentation import file_markup, plan_markup
    rendered = file_markup('<script>alert(1)</script>.txt', {'kind': 'document', 'rows': 1, 'columns': ['text']})
    assert '<script>' not in rendered and '&lt;script&gt;' in rendered
    payload = invoice_plan().model_dump()
    payload['steps'][2]['column'] = '<img src=x onerror=alert(1)>'
    rendered = plan_markup(Workflow.model_validate(payload))
    assert '<img' not in rendered and '&lt;img' in rendered


def test_strict_schema_uses_closed_objects_and_enum_constants():
    from prompt_compiler.provider import strict_schema
    from prompt_compiler.schema import CompilerResponse
    schema = strict_schema(CompilerResponse.model_json_schema())
    assert schema['additionalProperties'] is False
    assert schema['required'] == ['workflow', 'clarification']
    assert schema['$defs']['Read']['properties']['operation']['enum'] == ['read']
    assert 'title' not in schema['$defs']['Read']['properties']['id']
    assert 'title' in schema['$defs']['Workflow']['properties']
    assert 'title' in schema['$defs']['Workflow']['required']
