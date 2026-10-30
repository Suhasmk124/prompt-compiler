from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import re
import time
import operator
import pandas as pd
from openpyxl import load_workbook
from .files import inventory, read_file
from .schema import Workflow, Read, Extract, Filter, Flag, FilterRows, FlagRows, Calculate, Aggregate, Select, Sort, Export
from .validation import validate_workflow


@dataclass
class Run:
    tables: dict = field(default_factory=dict)
    artifacts: dict[str, bytes] = field(default_factory=dict)
    trace: list[dict] = field(default_factory=list)
    checks: list[str] = field(default_factory=list)
    error: str | None = None


def extract_invoice(row):
    text = str(row['text'])
    vendor = re.search(r'(?im)^\s*(?:vendor|supplier)\s*:\s*(.+)$', text)
    totals = re.findall(r'(?im)^\s*(?:grand total|invoice total|amount due|total)\s*:\s*(?:INR|Rs\.?|₹|USD|EUR|\$|€)?\s*([0-9][0-9,]*(?:\.[0-9]{1,2})?)\s*(?:INR|USD|EUR)?\s*$', text)
    currencies = set(re.findall(r'\b(?:INR|USD|EUR)\b', text.upper()))
    for symbol, currency in [('₹', 'INR'), ('$', 'USD'), ('€', 'EUR')]:
        if symbol in text:
            currencies.add(currency)
    if re.search(r'\bRs\.?\s', text):
        currencies.add('INR')
    valid = bool(vendor) and len(totals) == 1 and len(currencies) == 1
    return {'source_file': row['source_file'], 'vendor': vendor.group(1).strip() if vendor else None,
            'total': float(totals[0].replace(',', '')) if len(totals) == 1 else None,
            'currency': next(iter(currencies)) if len(currencies) == 1 else None,
            'extraction_status': 'ok' if valid else 'needs_review'}


def numeric(series):
    values = pd.to_numeric(series, errors='coerce')
    if (series.notna() & values.isna()).any():
        raise ValueError(f'Column {series.name} contains non-numeric values; clean the input first.')
    return values


def predicate(frame, step):
    values = frame[step.column]
    if isinstance(step.value, (float, int)) and not isinstance(step.value, bool):
        values = numeric(values)
    target = step.value
    if isinstance(target, str) and step.comparison in {'eq', 'ne'} and not step.case_sensitive:
        # Match natural-language values such as "upi" against "UPI" without
        # altering the actual source data or converting missing values to text.
        values = values.astype('string').str.strip().str.casefold()
        target = target.strip().casefold()
    if step.column == 'total' and 'currency' in frame:
        if frame['currency'].isna().any() or set(frame['currency'].dropna()) - {'INR'}:
            raise ValueError('Invoice thresholds currently require explicit INR currency on every row.')
    comparisons = {'gt': operator.gt, 'gte': operator.ge, 'lt': operator.lt,
                   'lte': operator.le, 'eq': operator.eq, 'ne': operator.ne}
    result = comparisons[step.comparison](values, target).astype('boolean')
    return result.mask(values.isna(), pd.NA)


def compound_predicate(frame, step):
    results = [predicate(frame, condition) for condition in step.conditions]
    combined = results[0]
    for condition in results[1:]:
        # Nullable Boolean logic: false AND unknown is false; true OR unknown
        # is true. Otherwise unknown remains explicit.
        combined = combined & condition if step.match == 'all' else combined | condition
    return combined


def export_table(frame, filename):
    # Escape spreadsheet formulas in both headers and data; never emit executable cells.
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
        return value
    clean = frame.copy().map(safe)
    clean.columns = [safe(str(c)) for c in clean.columns]
    if filename.endswith('.csv'):
        payload = clean.to_csv(index=False).encode('utf-8-sig')
        restored = pd.read_csv(BytesIO(payload))
        if restored.shape != clean.shape or list(restored.columns) != list(clean.columns):
            raise ValueError('CSV round-trip verification failed.')
        return payload
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        clean.to_excel(writer, sheet_name='Report', index=False)
        sheet = writer.sheets['Report']
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(45, max(14, max(len(str(c.value or '')) for c in col[:100]) + 2))
    payload = buffer.getvalue()
    book = load_workbook(BytesIO(payload), data_only=False)
    sheet = book.active
    if sheet.max_row != len(clean) + 1 or sheet.max_column != len(clean.columns):
        raise ValueError('Workbook dimensions do not match the result table.')
    if any(c.data_type == 'f' for row in sheet for c in row):
        raise ValueError('Unexpected formula in exported workbook.')
    return payload


def execute(workflow: Workflow, files: dict[str, bytes]) -> Run:
    manifest = inventory(files)
    validate_workflow(workflow, manifest)
    run = Run()
    for step in workflow.steps:
        start = time.perf_counter()
        try:
            if isinstance(step, Read):
                frame = pd.concat([read_file(n, files[n]) for n in step.files], ignore_index=True)
            else:
                frame = run.tables[step.source].copy()
                if isinstance(step, Extract):
                    frame = pd.DataFrame([extract_invoice(row) for _, row in frame.iterrows()])
                elif isinstance(step, (Filter, FilterRows)):
                    condition = compound_predicate(frame, step) if isinstance(step, FilterRows) else predicate(frame, step)
                    if condition.isna().any():
                        raise ValueError('Filter encountered missing values; resolve them before excluding rows.')
                    frame = frame.loc[condition].reset_index(drop=True)
                elif isinstance(step, (Flag, FlagRows)):
                    frame[step.output_column] = compound_predicate(frame, step) if isinstance(step, FlagRows) else predicate(frame, step)
                elif isinstance(step, Calculate):
                    left = numeric(frame[step.left])
                    right = numeric(frame[step.right]) if isinstance(step.right, str) else step.right
                    if step.arithmetic == 'divide' and ((right == 0).any() if isinstance(right, pd.Series) else right == 0):
                        raise ValueError('Division by zero is not allowed.')
                    frame[step.output_column] = {'add': operator.add, 'subtract': operator.sub,
                                                'multiply': operator.mul, 'divide': operator.truediv}[step.arithmetic](left, right)
                elif isinstance(step, Aggregate):
                    if step.function != 'count':
                        frame[step.column] = numeric(frame[step.column])
                    if frame[step.column].isna().any():
                        raise ValueError('Aggregation encountered missing values; resolve them first.')
                    if step.column == 'total' and 'currency' in frame:
                        if frame['currency'].isna().any() or frame['currency'].nunique() > 1:
                            raise ValueError('Cannot aggregate mixed or unknown invoice currencies.')
                    if step.group_by:
                        frame = frame.groupby(step.group_by, dropna=False)[step.column].agg(step.function).reset_index(name=step.output_column)
                    else:
                        value = getattr(frame[step.column], step.function)()
                        frame = pd.DataFrame([{step.output_column: value}])
                elif isinstance(step, Select):
                    frame = frame[step.columns]
                elif isinstance(step, Sort):
                    frame = frame.sort_values(step.column, ascending=not step.descending, kind='stable').reset_index(drop=True)
                elif isinstance(step, Export):
                    run.artifacts[step.filename] = export_table(frame, step.filename)
                    run.checks.append(f'{step.filename}: serialization and dimensions verified ({len(frame)} rows).')
            if len(frame) > 100_000:
                raise ValueError('Combined result exceeds 100,000 rows.')
            run.tables[step.id] = frame
            run.trace.append({'node': step.id, 'operation': step.operation, 'status': 'success',
                              'rows': len(frame), 'seconds': round(time.perf_counter() - start, 4)})
        except Exception as exc:
            run.error = f'{step.id}: {exc}'
            run.artifacts.clear()
            run.trace.append({'node': step.id, 'operation': step.operation, 'status': 'failed', 'error': str(exc),
                              'seconds': round(time.perf_counter() - start, 4)})
            break
    return run
