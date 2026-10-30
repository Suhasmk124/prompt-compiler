"""Local execution only: no model, provider stub, HTTP, or UI calls."""
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from prompt_compiler.compiler import validate_requested_actions
from prompt_compiler.executor import execute
from prompt_compiler.files import inventory
from prompt_compiler.schema import Workflow
from prompt_compiler.validation import validate_workflow
from prompt_compiler.workspace_tools import update_parameters


FILES = {'input.csv': b'invoice_id,payment_method,delivery_city\nA,UPI,Hyderabad\nB,upi,Pune\nC,NEFT,hyderabad\nD, UPI , Hyderabad \nE,,Hyderabad\nF,,Pune\n'}


def plan(operation='flag_rows', match='all', case_sensitive=False):
    rule = {'id': 'match_rows', 'label': 'Match payment and city', 'operation': operation,
            'source': 'read', 'match': match, 'conditions': [
                {'column': 'payment_method', 'comparison': 'eq', 'value': 'upi', 'case_sensitive': case_sensitive},
                {'column': 'delivery_city', 'comparison': 'eq', 'value': 'hyderabad', 'case_sensitive': case_sensitive}]}
    if operation == 'flag_rows':
        rule['output_column'] = 'matches_request'
    return Workflow.model_validate({'version': '1.0', 'title': 'Payment and city', 'steps': [
        {'id': 'read', 'label': 'Read', 'operation': 'read', 'files': ['input.csv']}, rule,
        {'id': 'report', 'label': 'Report', 'operation': 'export', 'source': 'match_rows', 'filename': 'report.xlsx'}]})


def test_and_flag_preserves_rows_and_exports_combined_boolean():
    result = execute(plan(), FILES)
    assert result.error is None
    table = result.tables['report']
    assert table['invoice_id'].tolist() == ['A', 'B', 'C', 'D', 'E', 'F']
    assert table['matches_request'].iloc[:4].tolist() == [True, False, False, True]
    assert pd.isna(table['matches_request'].iloc[4])
    assert table['matches_request'].iloc[5] == False
    assert table['payment_method'].iloc[3] == ' UPI '
    sheet = load_workbook(BytesIO(result.artifacts['report.xlsx'])).active
    assert sheet.max_column == 4 and sheet.max_row == 7
    assert [sheet.cell(r, 4).value for r in range(2, 8)] == [True, False, False, True, None, False]


def test_or_flag_uses_nullable_boolean_logic():
    result = execute(plan(match='any'), FILES)
    assert result.error is None
    assert result.tables['report']['matches_request'].iloc[:5].tolist() == [True] * 5
    assert pd.isna(result.tables['report']['matches_request'].iloc[5])


def test_explicit_case_sensitive_flag_does_not_fold_text():
    result = execute(plan(case_sensitive=True), FILES)
    assert result.error is None
    assert result.tables['report']['matches_request'].dropna().eq(False).all()


def test_compound_filter_stops_on_unknown_and_filters_when_complete():
    result = execute(plan(operation='filter_rows'), FILES)
    assert 'missing values' in result.error and not result.artifacts
    complete = {'input.csv': FILES['input.csv'].split(b'E,')[0]}
    result = execute(plan(operation='filter_rows'), complete)
    assert result.error is None
    assert result.tables['report']['invoice_id'].tolist() == ['A', 'D']


def test_compound_unknown_columns_and_output_overwrites_rejected():
    payload = plan().model_dump()
    payload['steps'][1]['conditions'][1]['column'] = 'imaginary_city'
    with pytest.raises(ValueError, match='unknown columns'):
        validate_workflow(Workflow.model_validate(payload), inventory(FILES))
    payload = plan().model_dump()
    payload['steps'][1]['output_column'] = 'invoice_id'
    with pytest.raises(ValueError, match='new, non-empty'):
        validate_workflow(Workflow.model_validate(payload), inventory(FILES))


def test_rule_edits_preserve_columns_and_can_change_match_logic():
    original = plan()
    revised = update_parameters(original, {'match_rows': {'match': 'any'}}, inventory(FILES))
    assert original.steps[1].match == 'all' and revised.steps[1].match == 'any'
    conditions = [c.model_dump() for c in original.steps[1].conditions]
    conditions[0]['column'] = 'invoice_id'
    with pytest.raises(ValueError, match='cannot change source columns'):
        update_parameters(original, {'match_rows': {'conditions': conditions}}, inventory(FILES))


def test_action_guard_rejects_ignored_or_discarded_flag():
    validate_requested_actions(plan(), 'FLAG invoices whose payment mode is upi and delivery city is hyderabad')
    payload = plan().model_dump()
    payload['steps'].pop(1)
    payload['steps'][-1]['source'] = 'read'
    with pytest.raises(ValueError, match='omits that action|no exported report'):
        validate_requested_actions(Workflow.model_validate(payload), 'Flag matching invoices')
    validate_requested_actions(Workflow.model_validate(payload), 'Export all rows unchanged')
    payload = plan().model_dump()
    payload['steps'].insert(2, {'id': 'select', 'label': 'Select original', 'operation': 'select',
                              'source': 'match_rows', 'columns': ['invoice_id']})
    payload['steps'][-1]['source'] = 'select'
    with pytest.raises(ValueError, match='omits that action|no exported report'):
        validate_requested_actions(Workflow.model_validate(payload), 'Flag matching invoices')


def test_single_condition_flags_also_match_text_case_insensitively():
    payload = plan().model_dump()
    payload['steps'][1] = {'id': 'match_rows', 'label': 'UPI', 'operation': 'flag', 'source': 'read',
                           'column': 'payment_method', 'comparison': 'eq', 'value': 'upi',
                           'output_column': 'matches_request'}
    result = execute(Workflow.model_validate(payload), FILES)
    assert result.error is None
    assert result.tables['report']['matches_request'].iloc[:4].tolist() == [True, True, False, True]
